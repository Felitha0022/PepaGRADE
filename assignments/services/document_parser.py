import os

from docx import Document
import pymupdf


def extract_text_from_docx(file_path):
    """
    Extract text from a Microsoft Word DOCX document.

    Reads:
    - Normal paragraphs
    - Text inside tables
    """

    document = Document(file_path)

    text = []

    # ---------------------------------------------------------
    # READ NORMAL PARAGRAPHS
    # ---------------------------------------------------------

    for paragraph in document.paragraphs:

        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text.append(paragraph_text)

    # ---------------------------------------------------------
    # READ TABLES
    # ---------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text.append(" | ".join(row_text))

    # ---------------------------------------------------------
    # COMBINE ALL TEXT
    # ---------------------------------------------------------

    extracted_text = "\n".join(text).strip()

    # Make sure the document actually contains text
    if not extracted_text:

        raise ValueError(
            "No readable text was found in the Word document. "
            "Please make sure the uploaded document contains text."
        )

    return extracted_text


def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF document.
    """

    doc = pymupdf.open(file_path)

    text = []

    for page in doc:

        page_text = page.get_text().strip()

        if page_text:
            text.append(page_text)

    doc.close()

    extracted_text = "\n".join(text).strip()

    if not extracted_text:

        raise ValueError(
            "No readable text was found in the PDF document. "
            "Please make sure the PDF contains selectable text."
        )

    return extracted_text


def extract_text(file_path):
    """
    Detect the file type and extract its text.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    elif extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file type. "
            "Please upload a DOCX or PDF file."
        )